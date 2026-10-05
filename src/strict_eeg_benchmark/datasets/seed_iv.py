"""SEED-IV preprocessed raw EEG with provider-verified session labels."""

from __future__ import annotations

import re
from collections.abc import Iterator
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from scipy.io import loadmat

from ..features import DifferentialEntropyExtractor
from ..types import EEGTrial, FeatureBatch
from .base import EEGDataset

_FILE = re.compile(r"^(?P<subject>\d+)_(?P<date>\d{8})\.mat$")
_TRIAL = re.compile(r".+_eeg(?P<number>\d+)$")
_LABEL = re.compile(r"session(?P<session>[123])_label\s*=\s*\[(?P<values>[^]]+)\]", re.IGNORECASE)


class SEEDIVDataset(EEGDataset):
    name = "SEED-IV"
    sampling_rate = 200.0
    class_names = {0: "neutral", 1: "sad", 2: "fear", 3: "happy"}

    def __init__(self, root: str | Path):
        self.root = Path(root).expanduser().resolve()
        self.raw_dir = self.root / "eeg_raw_data"
        self.feature_dir = self.root / "eeg_feature_smooth"
        self._labels = self._read_labels()
        self._channel_names = self._read_channels()
        self._raw_files = self._discover(self.raw_dir)
        self._feature_files = self._discover(self.feature_dir)
        if set(self._raw_files) != set(self._feature_files):
            raise ValueError("SEED-IV raw and supplied-feature subject/session files differ")
        if set(self._raw_files) != {(str(subject), str(session))
                                   for subject in range(1, 16) for session in range(1, 4)}:
            raise ValueError("Expected exactly 15 subjects × 3 sessions")
        self._verify_stimulation_labels()

    def _read_labels(self) -> dict[str, tuple[int, ...]]:
        readme = (self.root / "ReadMe.txt").read_text(encoding="utf-8", errors="replace")
        parsed = {}
        for match in _LABEL.finditer(readme):
            values = tuple(int(value) for value in re.findall(r"\d+", match.group("values")))
            if len(values) != 24 or set(values) != {0, 1, 2, 3}:
                raise ValueError(f"Unexpected SEED-IV session {match.group('session')} labels")
            if any(values.count(label) != 6 for label in range(4)):
                raise ValueError("Expected six trials of each emotion per session")
            parsed[match.group("session")] = values
        if set(parsed) != {"1", "2", "3"}:
            raise ValueError("Provider README does not define all three session label lists")
        return parsed

    def _read_channels(self) -> tuple[str, ...]:
        workbook = load_workbook(self.root / "Channel Order.xlsx", read_only=True, data_only=True)
        sheet = workbook[workbook.sheetnames[0]]
        names = tuple(str(row[0]).strip().upper() for row in sheet.iter_rows(values_only=True) if row[0])
        workbook.close()
        if len(names) != 62 or len(set(names)) != 62:
            raise ValueError("Expected 62 unique provider channel names")
        return names

    @staticmethod
    def _discover(directory: Path) -> dict[tuple[str, str], Path]:
        found = {}
        for session in ("1", "2", "3"):
            for path in (directory / session).glob("*.mat"):
                match = _FILE.fullmatch(path.name)
                if not match:
                    raise ValueError(f"Unexpected SEED-IV MATLAB filename: {path}")
                key = (match.group("subject"), session)
                if key in found:
                    raise ValueError(f"Duplicate SEED-IV subject/session: {key}")
                found[key] = path
        return found

    def _verify_stimulation_labels(self) -> None:
        workbook = load_workbook(self.root / "SEED-IV_stimulation.xlsx", read_only=True, data_only=True)
        if len(workbook.worksheets) < 3:
            raise ValueError("Stimulation workbook lacks three session sheets")
        for session in ("1", "2", "3"):
            sheet = workbook.worksheets[int(session) - 1]
            observed = tuple(int(sheet.cell(row=i + 2, column=2).value) for i in range(24))
            if observed != self._labels[session]:
                raise AssertionError(f"Session {session} stimulation workbook labels differ from README")
        workbook.close()

    @property
    def channel_names(self) -> tuple[str, ...]:
        return self._channel_names

    @property
    def subject_ids(self) -> tuple[str, ...]:
        return tuple(str(i) for i in range(1, 16))

    @property
    def subject_session_files(self) -> tuple[tuple[str, str, Path, Path], ...]:
        return tuple((subject, session, self._raw_files[(subject, session)], self._feature_files[(subject, session)])
                     for subject in self.subject_ids for session in ("1", "2", "3"))

    @property
    def labels_by_session(self) -> dict[str, tuple[int, ...]]:
        return self._labels.copy()

    def iter_trials(self, subjects: set[str] | None = None) -> Iterator[EEGTrial]:
        wanted = None if subjects is None else set(map(str, subjects))
        for subject, session, path, _ in self.subject_session_files:
            if wanted is not None and subject not in wanted:
                continue
            mat = loadmat(path)
            arrays = {int(match.group("number")): np.asarray(value)
                      for name, value in mat.items() if (match := _TRIAL.fullmatch(name))}
            if set(arrays) != set(range(1, 25)):
                raise ValueError(f"Expected trials 1..24 in {path}, found {sorted(arrays)}")
            for number in range(1, 25):
                eeg = arrays[number]
                if eeg.ndim != 2 or eeg.shape[0] != 62 or eeg.shape[1] < 800:
                    raise ValueError(f"Invalid raw SEED-IV EEG shape {eeg.shape} in {path} trial {number}")
                yield EEGTrial(eeg=eeg, label=self._labels[session][number - 1], subject_id=subject,
                               session_id=session, trial_id=f"{subject}_session_{session}_trial_{number:02d}",
                               channel_names=self.channel_names, sampling_rate=self.sampling_rate,
                               dataset_name=self.name)

    def load_de_features(self, subjects: set[str] | None = None,
                         max_windows_per_trial: int | None = None) -> FeatureBatch:
        """Optional stateless DE extraction; replication models use raw EEG only."""
        if max_windows_per_trial is not None and max_windows_per_trial < 1:
            raise ValueError("max_windows_per_trial must be positive")
        extractor = DifferentialEntropyExtractor(self.sampling_rate, window_seconds=1.0)
        parts, labels, subject_ids, session_ids, trial_ids = [], [], [], [], []
        for trial in self.iter_trials(subjects):
            eeg = trial.eeg if max_windows_per_trial is None else trial.eeg[:, :max_windows_per_trial * 200]
            windows = extractor.transform_trial(eeg).reshape(-1, 310)
            n = len(windows)
            parts.append(windows)
            labels.append(np.full(n, trial.label, dtype=np.int64))
            subject_ids.append(np.full(n, trial.subject_id, dtype=object))
            session_ids.append(np.full(n, trial.session_id, dtype=object))
            trial_ids.append(np.full(n, trial.trial_id, dtype=object))
        if not parts:
            raise ValueError("No SEED-IV trials matched the requested subjects")
        feature_names = tuple(f"{channel}_{band}" for channel in self.channel_names
                              for band in ("delta", "theta", "alpha", "beta", "gamma"))
        return FeatureBatch(X=np.concatenate(parts), y=np.concatenate(labels),
                            subject_ids=np.concatenate(subject_ids), session_ids=np.concatenate(session_ids),
                            trial_ids=np.concatenate(trial_ids), dataset_name=self.name,
                            feature_names=feature_names)
