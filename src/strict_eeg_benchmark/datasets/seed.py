from __future__ import annotations

import re
from collections.abc import Iterator
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from scipy.io import loadmat

from ..features import DifferentialEntropyExtractor
from ..types import EEGTrial, FeatureBatch
from .base import EEGDataset


_FILE_RE = re.compile(r"^(?P<subject>\d+)_(?P<date>\d{8})\.mat$")
_EEG_KEY_RE = re.compile(r".+_eeg(?P<trial>\d+)$")
_DE_KEY_RE = re.compile(r"de_movingAve(?P<trial>\d+)$")


class SEEDDataset(EEGDataset):
    name = "SEED"
    sampling_rate = 200.0
    label_map = {-1: 0, 0: 1, 1: 2}
    class_names = {0: "negative", 1: "neutral", 2: "positive"}
    band_names = ("delta", "theta", "alpha", "beta", "gamma")

    def __init__(self, root: str | Path):
        self.root = Path(root).expanduser().resolve()
        self.preprocessed_dir = self.root / "Preprocessed_EEG"
        self.features_dir = self.root / "ExtractedFeatures_1s"
        for required in (self.preprocessed_dir, self.features_dir):
            if not required.is_dir():
                raise FileNotFoundError(f"Required SEED directory not found: {required}")
        self._labels = self._load_labels()
        self._channel_names = self._load_channel_names()
        self._raw_files = self._discover(self.preprocessed_dir)
        self._feature_files = self._discover(self.features_dir)
        if set(self._raw_files) != set(self._feature_files):
            raise ValueError("SEED raw/preprocessed and feature subject-session files do not match")

    @staticmethod
    def _discover(directory: Path) -> dict[tuple[str, str], Path]:
        found: dict[tuple[str, str], Path] = {}
        for path in directory.glob("*.mat"):
            match = _FILE_RE.match(path.name)
            if match:
                key = (match.group("subject"), match.group("date"))
                if key in found:
                    raise ValueError(f"Duplicate SEED subject-session file: {key}")
                found[key] = path
        if not found:
            raise FileNotFoundError(f"No SEED subject-session .mat files found in {directory}")
        return found

    def _load_labels(self) -> np.ndarray:
        path = self.preprocessed_dir / "label.mat"
        labels = np.asarray(loadmat(path)["label"]).reshape(-1)
        if len(labels) != 15 or not set(map(int, labels)).issubset(self.label_map):
            raise ValueError(f"Unexpected SEED labels in {path}: {labels.tolist()}")
        return np.asarray([self.label_map[int(value)] for value in labels], dtype=np.int64)

    def _load_channel_names(self) -> tuple[str, ...]:
        path = self.root / "channel-order.xlsx"
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook[workbook.sheetnames[0]]
        names = tuple(str(row[0]).strip().upper() for row in sheet.iter_rows(values_only=True) if row[0])
        workbook.close()
        if len(names) != 62 or len(set(names)) != 62:
            raise ValueError(f"Expected 62 unique SEED channels in {path}, found {len(names)}")
        return names

    @property
    def channel_names(self) -> tuple[str, ...]:
        return self._channel_names

    @property
    def subject_ids(self) -> tuple[str, ...]:
        return tuple(sorted({key[0] for key in self._raw_files}, key=int))

    @property
    def subject_session_files(self) -> tuple[tuple[str, str, Path, Path], ...]:
        keys = sorted(self._raw_files, key=lambda item: (int(item[0]), item[1]))
        return tuple((s, d, self._raw_files[(s, d)], self._feature_files[(s, d)]) for s, d in keys)

    def iter_trials(self, subjects: set[str] | None = None) -> Iterator[EEGTrial]:
        wanted = None if subjects is None else set(map(str, subjects))
        for subject, session, raw_path, _ in self.subject_session_files:
            if wanted is not None and subject not in wanted:
                continue
            mat = loadmat(raw_path)
            trial_arrays: dict[int, np.ndarray] = {}
            for key, value in mat.items():
                match = _EEG_KEY_RE.match(key)
                if match:
                    trial_arrays[int(match.group("trial"))] = np.asarray(value)
            if set(trial_arrays) != set(range(1, 16)):
                raise ValueError(f"Expected trials 1..15 in {raw_path}, found {sorted(trial_arrays)}")
            for trial_number in range(1, 16):
                eeg = trial_arrays[trial_number]
                if eeg.ndim != 2 or eeg.shape[0] != len(self.channel_names):
                    raise ValueError(f"Unexpected EEG shape {eeg.shape} in {raw_path}, trial {trial_number}")
                yield EEGTrial(
                    eeg=eeg,
                    label=int(self._labels[trial_number - 1]),
                    subject_id=subject,
                    session_id=session,
                    trial_id=f"{subject}_{session}_trial_{trial_number:02d}",
                    channel_names=self.channel_names,
                    sampling_rate=self.sampling_rate,
                    dataset_name=self.name,
                )

    def load_de_features(
        self, subjects: set[str] | None = None, max_windows_per_trial: int | None = None
    ) -> FeatureBatch:
        """Extract stateless 1-second DE windows directly from preprocessed EEG."""
        if max_windows_per_trial is not None and max_windows_per_trial < 1:
            raise ValueError("max_windows_per_trial must be positive")
        extractor = DifferentialEntropyExtractor(self.sampling_rate, window_seconds=1.0)
        xs: list[np.ndarray] = []
        ys: list[np.ndarray] = []
        subject_ids: list[np.ndarray] = []
        session_ids: list[np.ndarray] = []
        trial_ids: list[np.ndarray] = []
        for trial in self.iter_trials(subjects):
            eeg = trial.eeg
            if max_windows_per_trial is not None:
                eeg = eeg[:, : max_windows_per_trial * extractor.window_samples]
            windows = extractor.transform_trial(eeg)
            flat = windows.reshape(len(windows), -1)
            n = len(flat)
            xs.append(flat)
            ys.append(np.full(n, trial.label, dtype=np.int64))
            subject_ids.append(np.full(n, trial.subject_id, dtype=object))
            session_ids.append(np.full(n, trial.session_id, dtype=object))
            trial_ids.append(np.full(n, trial.trial_id, dtype=object))
        return self._combine_feature_parts(xs, ys, subject_ids, session_ids, trial_ids)

    def load_provider_de_features(
        self, subjects: set[str] | None = None, max_windows_per_trial: int | None = None
    ) -> FeatureBatch:
        """Load the provider's temporally smoothed DE export with provenance."""
        if max_windows_per_trial is not None and max_windows_per_trial < 1:
            raise ValueError("max_windows_per_trial must be positive")
        wanted = None if subjects is None else set(map(str, subjects))
        xs: list[np.ndarray] = []
        ys: list[np.ndarray] = []
        subject_ids: list[np.ndarray] = []
        session_ids: list[np.ndarray] = []
        trial_ids: list[np.ndarray] = []

        for subject, session, _, feature_path in self.subject_session_files:
            if wanted is not None and subject not in wanted:
                continue
            mat = loadmat(feature_path)
            for trial_number in range(1, 16):
                key = f"de_movingAve{trial_number}"
                if key not in mat:
                    raise ValueError(f"Missing {key} in {feature_path}")
                values = np.asarray(mat[key], dtype=np.float64)
                if values.ndim != 3 or values.shape[0] != 62 or values.shape[2] != 5:
                    raise ValueError(f"Unexpected {key} shape {values.shape} in {feature_path}")
                windows = np.transpose(values, (1, 0, 2))
                if max_windows_per_trial is not None:
                    windows = windows[:max_windows_per_trial]
                flat = windows.reshape(len(windows), -1)
                n = len(flat)
                xs.append(flat)
                ys.append(np.full(n, self._labels[trial_number - 1], dtype=np.int64))
                subject_ids.append(np.full(n, subject, dtype=object))
                session_ids.append(np.full(n, session, dtype=object))
                trial_ids.append(np.full(n, f"{subject}_{session}_trial_{trial_number:02d}", dtype=object))

        return self._combine_feature_parts(xs, ys, subject_ids, session_ids, trial_ids)

    def _combine_feature_parts(
        self,
        xs: list[np.ndarray],
        ys: list[np.ndarray],
        subject_ids: list[np.ndarray],
        session_ids: list[np.ndarray],
        trial_ids: list[np.ndarray],
    ) -> FeatureBatch:
        if not xs:
            raise ValueError("No SEED feature windows matched the requested subjects")
        feature_names = tuple(f"{channel}_{band}" for channel in self.channel_names for band in self.band_names)
        return FeatureBatch(
            X=np.concatenate(xs, axis=0),
            y=np.concatenate(ys),
            subject_ids=np.concatenate(subject_ids),
            session_ids=np.concatenate(session_ids),
            trial_ids=np.concatenate(trial_ids),
            dataset_name=self.name,
            feature_names=feature_names,
        )
