"""Prepare local controlled-study inputs directly from licensed provider files.

No downloads, raw-file edits, or model fitting are performed. Existing output
files are not overwritten. The generated arrays and metadata stay local.
"""
from __future__ import annotations

import argparse
import csv
import json
import pickle
import re
import sys
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from scipy.io import loadmat

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))


def assert_new(*paths: Path) -> None:
    if any(p.exists() for p in paths):
        raise FileExistsError("A prepared output already exists; use a separate output directory for validation")
    for p in paths:
        p.parent.mkdir(parents=True, exist_ok=True)


def prepare_seed(raw: Path, output: Path) -> None:
    a, b = output / "seed_lds_cache.npz", output / "seed_annotated_cache.npz"
    assert_new(a, b)
    labels = np.asarray(loadmat(raw / "Preprocessed_EEG/label.mat")["label"]).ravel().astype(int) + 1
    if labels.shape != (15,) or set(labels) != {0, 1, 2}:
        raise AssertionError("Unexpected provider SEED class map")
    pattern = re.compile(r"^(\d+)_(\d{8})\.mat$")
    files = [p for p in (raw / "ExtractedFeatures_1s").glob("*.mat") if pattern.match(p.name)]
    files.sort(key=lambda p: (int(pattern.match(p.name)[1]), pattern.match(p.name)[2]))
    if len(files) != 45:
        raise AssertionError("Expected 45 SEED subject-session feature files")
    x, y, subject, session, trial = [], [], [], [], []
    for path in files:
        sid, date = pattern.match(path.name).groups()
        data = loadmat(path)
        for number in range(1, 16):
            values = np.asarray(data[f"de_LDS{number}"], dtype=np.float32).transpose(1, 0, 2)
            if values.shape[1:] != (62, 5) or not np.isfinite(values).all():
                raise AssertionError("Invalid SEED provider feature array")
            n = len(values)
            x.append(values)
            y.extend([int(labels[number - 1])] * n)
            subject.extend([sid] * n)
            session.extend([date] * n)
            trial.extend([f"{sid}_{date}_trial_{number:02d}"] * n)
    arrays = {"data": np.concatenate(x), "labels": np.array(y, dtype=np.int64),
              "subjects": np.array(subject, dtype="U2"), "sessions": np.array(session, dtype="U8")}
    if arrays["data"].shape != (152730, 62, 5) or len(set(trial)) != 675:
        raise AssertionError("SEED provider totals changed")
    np.savez_compressed(a, **arrays)
    np.savez_compressed(b, labels=arrays["labels"], subjects=arrays["subjects"],
                        sessions=arrays["sessions"], trials=np.array(trial))
    print(f"SEED prepared: {len(y)} windows, 675 trials", flush=True)


def prepare_seediv(raw: Path, output: Path) -> None:
    from src.strict_eeg_benchmark.datasets.seed_iv import SEEDIVDataset
    assert_new(output, output.with_suffix(".json"))
    dataset = SEEDIVDataset(raw)
    x, y, subject, session, trial = [], [], [], [], []
    for sid, visit, _, feature in dataset.subject_session_files:
        data = loadmat(feature)
        for number in range(1, 25):
            values = np.asarray(data[f"de_LDS{number}"], dtype=np.float32).transpose(1, 0, 2)
            if values.shape[1:] != (62, 5) or not np.isfinite(values).all():
                raise AssertionError("Invalid SEED-IV provider feature array")
            n = len(values)
            x.append(values)
            y.extend([dataset.labels_by_session[visit][number - 1]] * n)
            subject.extend([sid] * n)
            session.extend([visit] * n)
            trial.extend([f"{sid}_{visit}_{number:02d}"] * n)
    arrays = {"x": np.concatenate(x), "y": np.array(y, dtype=np.int8),
              "subject": np.array(subject, dtype="U2"), "session": np.array(session, dtype="U1"),
              "trial": np.array(trial, dtype="U10")}
    if arrays["x"].shape != (37575, 62, 5) or len(set(trial)) != 1080:
        raise AssertionError("SEED-IV provider totals changed")
    np.savez_compressed(output, **arrays)
    output.with_suffix(".json").write_text(json.dumps({"windows": len(y), "trials": 1080,
        "representation": "provider de_LDS; 62 channels x 5 bands; no local population fitting"}, indent=2))
    print(f"SEED-IV prepared: {len(y)} windows, 1080 trials", flush=True)


def faced_labels(raw: Path) -> dict[int, int]:
    classes = ["anger", "disgust", "fear", "sadness", "neutral", "amusement", "inspiration", "joy", "tenderness"]
    book = load_workbook(raw / "Stimuli_info.xlsx", read_only=True, data_only=True)
    labels = {}
    for row in book.active.values:
        if isinstance(row[0], int) and 1 <= row[0] <= 28:
            emotion = str(row[5]).strip().lower()
            if emotion == "\\":
                emotion = "neutral"
            labels[row[0]] = classes.index(emotion)
    book.close()
    if set(labels) != set(range(1, 29)):
        raise AssertionError("FACED stimulus labels are incomplete")
    book = load_workbook(raw / "Task_event.xlsx", read_only=True, data_only=True)
    check = {int(row[0]): classes.index(str(row[1]).strip().lower())
             for row in book.active.values if isinstance(row[0], int) and 1 <= row[0] <= 28}
    book.close()
    if labels != check:
        raise AssertionError("FACED provider stimulus label sources disagree")
    return labels


def prepare_faced(raw: Path, output: Path, smoke_subjects: list[str] | None) -> None:
    from src.strict_eeg_benchmark.features import DifferentialEntropyExtractor
    assert_new(output)
    with (raw / "Recording_info.csv").open(newline="", encoding="utf-8-sig") as handle:
        people = sorted(r["sub"].strip() for r in csv.DictReader(handle))
    if people != [f"sub{i:03d}" for i in range(123)]:
        raise AssertionError("FACED recording identifiers changed")
    labels = faced_labels(raw)
    selected = smoke_subjects or people
    if not set(selected) <= set(people):
        raise ValueError("Unknown FACED smoke subject")
    selected = sorted(selected)
    extractor = DifferentialEntropyExtractor(250, window_seconds=1)
    x, y, subject, trial, local = [], [], [], [], []
    for number, sid in enumerate(selected, 1):
        with (raw / "Processed_data" / f"{sid}.pkl").open("rb") as handle:
            eeg = np.asarray(pickle.load(handle))
        if eeg.shape != (28, 32, 7500) or not np.isfinite(eeg).all():
            raise AssertionError("Invalid FACED processed subject array")
        # Match the archived raw-window cache's float32 boundary exactly before
        # the stateless extractor performs its float64 spectral calculation.
        windows = eeg.astype(np.float32).reshape(28, 32, 30, 250).transpose(0, 2, 1, 3).reshape(840, 32, 250)
        signal = windows.astype(np.float64).transpose(1, 0, 2).reshape(32, 840 * 250)
        values = extractor.transform_trial(signal).astype(np.float32)
        if values.shape != (840, 32, 5) or not np.isfinite(values).all():
            raise AssertionError("Invalid FACED DE features")
        x.append(values)
        y.extend(np.repeat([labels[v] for v in range(1, 29)], 30))
        subject.extend([sid] * 840)
        trial.extend(np.repeat([f"{sid}_video_{v:02d}" for v in range(1, 29)], 30))
        local.extend(np.tile(np.arange(30), 28))
        if number % 10 == 0 or number == len(selected):
            print(f"FACED prepared subjects {number}/{len(selected)}", flush=True)
    np.savez_compressed(output, x=np.concatenate(x), y=np.array(y, dtype=np.int64),
        subject=np.array(subject), trial=np.array(trial), local_window=np.array(local, dtype=np.int16),
        method_id=np.array("faced_stateless_welch_de_float32_boundary_v1"),
        status=np.array("complete" if selected == people else "smoke_only"))


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--dataset", choices=["SEED", "SEED-IV", "FACED"], required=True)
    parser.add_argument("--data-parent", type=Path, required=True,
                        help="Directory containing licensed data/SEED, data/SEED-IV, data/FACED")
    parser.add_argument("--project-root", type=Path, default=Path(__file__).resolve().parents[1])
    parser.add_argument("--seed-cache-dir", type=Path,
                        help="SEED output; defaults to the source data/SEED directory")
    parser.add_argument("--faced-smoke-subjects", nargs="+",
                        help="Prepare a separately named partial FACED check, not a model input")
    args = parser.parse_args()
    raw = args.data_parent / "data" / args.dataset
    if args.dataset == "SEED":
        prepare_seed(raw, args.seed_cache_dir or raw)
    elif args.dataset == "SEED-IV":
        prepare_seediv(raw, args.project_root / "experiments/p4_cross_dataset/seediv_provider_de_lds.npz")
    else:
        name = "faced_features_smoke.npz" if args.faced_smoke_subjects else "faced_features.npz"
        prepare_faced(raw, args.project_root / "experiments/controlled_protocol_revision" / name,
                      args.faced_smoke_subjects)


if __name__ == "__main__":
    main()
